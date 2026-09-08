from __future__ import annotations

from pathlib import Path
from typing import Any

from investigation_rag.extractors.base import BaseExtractor, ExtractedDocument, ExtractedSection
from investigation_rag.extractors.structure import SourceBlock


class DocxExtractor(BaseExtractor):
    source_type = "docx"

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required to extract .docx files") from exc

        document = Document(str(path))
        metadata = {"path": str(path), **_core_properties(document)}
        plain_parts: list[str] = []
        markdown_parts: list[str] = []
        source_blocks: list[SourceBlock] = []
        list_counts: dict[str, int] = {}
        for block_index, block in enumerate(_iter_blocks(document)):
            if _is_table(block):
                for row_index, row in enumerate(block.rows):
                    cells = []
                    row_parts = []
                    runs = []
                    offset = 0
                    for cell_index, cell in enumerate(row.cells):
                        cell_text = "\n".join(_source_paragraph_text(p) for p in cell.paragraphs)
                        cells.append({"column": cell_index, "start": offset, "end": offset + len(cell_text)})
                        inner_offset = offset
                        for paragraph in cell.paragraphs:
                            runs.extend(_paragraph_runs(paragraph, inner_offset))
                            inner_offset += len(_source_paragraph_text(paragraph)) + 1
                        row_parts.append(cell_text)
                        offset += len(cell_text) + 1
                    source_blocks.append(SourceBlock(
                        f"b{block_index}-r{row_index}", "\t".join(row_parts), "table_row",
                        {"table": block_index, "row": row_index, "cells": cells,
                         "is_header": bool(row._tr.xpath('./w:trPr/w:tblHeader'))}, runs,
                    ))
                plain_table, markdown_table = _table_text(block)
                if plain_table:
                    plain_parts.append(plain_table)
                    markdown_parts.append(markdown_table)
                continue

            source_text = _source_paragraph_text(block)
            text = source_text.strip()
            if not text:
                continue
            style = block.style.name.lower() if block.style else ""
            heading_level = _heading_level(style) if style.startswith("heading") else 1 if style == "title" else 0
            numbering = block._p.xpath('./w:pPr/w:numPr')
            is_list = bool(numbering) or "list bullet" in style or "list number" in style
            marker, list_level = _list_marker(block, document, list_counts, style)
            source_blocks.append(SourceBlock(
                f"b{block_index}", source_text,
                "heading" if heading_level else "list_item" if is_list else "paragraph",
                {"paragraph": block_index, "list_marker": marker, "list_level": list_level},
                _paragraph_runs(block), heading_level,
            ))
            plain_parts.append(text)
            markdown_parts.append(_paragraph_markdown(block))

        text = "\n\n".join(plain_parts).strip()
        markdown = "\n\n".join(markdown_parts).strip()
        return ExtractedDocument(
            source_type=self.source_type,
            sections=[ExtractedSection(text=text, source_type=self.source_type, display_markdown=markdown or None,
                                       blocks=source_blocks)],
            metadata=metadata,
        )


def _source_paragraph_text(paragraph: Any) -> str:
    from docx.text.run import Run
    return "".join(Run(element, paragraph).text for element in paragraph._p.xpath('.//w:r'))


def _list_marker(paragraph: Any, document: Any, counts: dict, style: str) -> tuple[str, int]:
    num_pr = paragraph._p.xpath('./w:pPr/w:numPr')
    if not num_pr and paragraph.style is not None:
        num_pr = paragraph.style.element.xpath('./w:pPr/w:numPr')
    level, number_format, pattern, start = 0, "bullet", "•", 1
    key = style
    if num_pr:
        from docx.oxml.ns import qn
        num_ids = num_pr[0].xpath('./w:numId')
        levels = num_pr[0].xpath('./w:ilvl')
        level = int(levels[0].get(qn('w:val'))) if levels else 0
        if num_ids:
            num_id = num_ids[0].get(qn('w:val'))
            key = str(num_id)
            root = document.part.numbering_part.element
            nums = root.xpath(f'./w:num[@w:numId="{num_id}"]/w:abstractNumId')
            if nums:
                abstract_id = nums[0].get(qn('w:val'))
                definitions = root.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{level}"]')
                if definitions:
                    definition = definitions[0]
                    formats = definition.findall(qn('w:numFmt'))
                    patterns = definition.findall(qn('w:lvlText'))
                    starts = definition.findall(qn('w:start'))
                    number_format = formats[0].get(qn('w:val')) if formats else "bullet"
                    pattern = patterns[0].get(qn('w:val')) if patterns else "%1."
                    start = int(starts[0].get(qn('w:val'))) if starts else 1
    elif "list number" in style:
        number_format, pattern = "decimal", "%1."
    if number_format == "bullet":
        return "•", level
    key_level = (key, level)
    counts[key_level] = counts.get(key_level, start - 1) + 1
    for deeper in list(counts):
        if isinstance(deeper, tuple) and deeper[0] == key and deeper[1] > level:
            del counts[deeper]
    for index in range(1, 10):
        value = counts.get((key, index - 1), 1)
        rendered = chr(96 + value) if number_format == "lowerLetter" and 1 <= value <= 26 else chr(64 + value) if number_format == "upperLetter" and 1 <= value <= 26 else str(value)
        pattern = pattern.replace(f"%{index}", rendered)
    return pattern, level


def _paragraph_runs(paragraph: Any, offset: int = 0) -> list[dict[str, Any]]:
    runs = []
    # Include hyperlink runs in their actual XML order too.
    from docx.text.run import Run
    for element in paragraph._p.xpath('.//w:r'):
        run = Run(element, paragraph)
        text = run.text
        runs.append({"start": offset, "end": offset + len(text),
                     "bold": bool(run.bold), "italic": bool(run.italic), "underline": bool(run.underline)})
        offset += len(text)
    return runs


def _core_properties(document: object) -> dict[str, object]:
    properties = getattr(document, "core_properties", None)
    if properties is None:
        return {}
    field_map = {
        "title": "title",
        "author": "author",
        "subject": "subject",
        "keywords": "keywords",
        "comments": "comments",
        "category": "category",
        "last_modified_by": "last_modified_by",
        "revision": "revision",
        "created": "creation_date",
        "modified": "modification_date",
    }
    metadata: dict[str, object] = {}
    for attr, key in field_map.items():
        value = getattr(properties, attr, None)
        if value in (None, ""):
            continue
        metadata[key] = value.isoformat() if hasattr(value, "isoformat") else str(value)
    return metadata


def _iter_blocks(document: Any) -> list[Any]:
    try:
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return list(getattr(document, "paragraphs", []))

    blocks = []
    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            blocks.append(Paragraph(child, document))
        elif isinstance(child, CT_Tbl):
            blocks.append(Table(child, document))
    return blocks


def _is_table(block: Any) -> bool:
    return hasattr(block, "rows") and hasattr(block, "columns")


def _paragraph_markdown(paragraph: Any) -> str:
    text = paragraph.text.strip()
    style = paragraph.style.name if paragraph.style else ""
    style_name = style.lower()
    inline = _inline_markdown(paragraph) or _escape_markdown(text)

    if style_name.startswith("heading"):
        level = _heading_level(style_name)
        return f"{'#' * level} {inline}"
    if "list bullet" in style_name:
        return f"- {inline}"
    if "list number" in style_name:
        return f"1. {inline}"
    return inline


def _heading_level(style_name: str) -> int:
    digits = "".join(char for char in style_name if char.isdigit())
    if not digits:
        return 1
    return min(max(int(digits), 1), 6)


def _inline_markdown(paragraph: Any) -> str:
    parts: list[str] = []
    for run in paragraph.runs:
        text = run.text
        if not text:
            continue
        escaped = _escape_markdown(text)
        if run.bold and run.italic:
            parts.append(f"***{escaped}***")
        elif run.bold:
            parts.append(f"**{escaped}**")
        elif run.italic:
            parts.append(f"*{escaped}*")
        else:
            parts.append(escaped)
    return "".join(parts).strip()


def _table_text(table: Any) -> tuple[str, str]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return "", ""

    plain = "\n".join("\t".join(cells) for cells in rows)
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    separator = ["---" for _ in range(width)]
    body = normalized[1:] or [["" for _ in range(width)]]
    markdown_rows = [header, separator, *body]
    markdown = "\n".join("| " + " | ".join(_escape_table_cell(cell) for cell in row) + " |" for row in markdown_rows)
    return plain, markdown


def _escape_markdown(text: str) -> str:
    return text.replace("\\", "\\\\").replace("*", "\\*").replace("_", "\\_").replace("`", "\\`")


def _escape_table_cell(text: str) -> str:
    return _escape_markdown(text).replace("|", "\\|")
