from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from investigation_rag.evaluation import (
    EVALUATION_ERROR_COLUMN,
    GENERATED_ANSWER_COLUMN,
    JUDGE_SCORE_COLUMN,
    SUMMARY_SHEET_NAME,
    EvaluationOutcome,
    EvaluationReportParameters,
    EvaluationRow,
    build_judge_prompt,
    export_evaluation_report,
    load_evaluation_workbook,
    parse_judge_score,
    summarize_evaluation,
)
from investigation_rag.gui.evaluation_panel import EvaluationWorker


def test_loads_supplied_evaluation_workbook() -> None:
    path = Path(__file__).parents[2] / "eval" / "eval_qs.xlsx"

    evaluation = load_evaluation_workbook(path)

    assert evaluation.sheet_name == "Blad1"
    assert [row.question_id for row in evaluation.rows] == ["q0011", "q0012"]
    assert [row.max_points for row in evaluation.rows] == [Decimal("2"), Decimal("1")]
    assert evaluation.rows[0].comment == "Elk correct antwoord is een punt"


def test_workbook_requires_question_answer_and_points(tmp_path: Path) -> None:
    path = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Question", "Answer"])
    worksheet.append(["What happened?", "Nothing."])
    workbook.save(path)

    with pytest.raises(ValueError, match="required columns"):
        load_evaluation_workbook(path)


def test_judge_prompt_contains_all_scoring_inputs() -> None:
    row = EvaluationRow(
        worksheet_row=2,
        question_id="q1",
        question="What happened?",
        ground_truth_answer="A meeting occurred.",
        comment="One point for mentioning the meeting.",
        max_points=Decimal("2.5"),
    )

    prompt = build_judge_prompt(row, "There was a meeting.")

    assert "What happened?" in prompt
    assert "A meeting occurred." in prompt
    assert "There was a meeting." in prompt
    assert "One point for mentioning the meeting." in prompt
    assert "2.5" in prompt
    assert "Return only the awarded number of points" in prompt


@pytest.mark.parametrize(
    ("response", "expected"),
    [("0", Decimal("0")), ("1", Decimal("1")), (" 1.5 ", Decimal("1.5")), ("1,5", Decimal("1.5"))],
)
def test_parse_judge_score_accepts_only_bounded_numbers(response: str, expected: Decimal) -> None:
    assert parse_judge_score(response, Decimal("2")) == expected


@pytest.mark.parametrize("response", ["Score: 1", "1 point", "-1", "3", "1\nReasoning", "NaN", ""])
def test_parse_judge_score_rejects_invalid_output(response: str) -> None:
    with pytest.raises(ValueError):
        parse_judge_score(response, Decimal("2"))


def test_judge_retries_malformed_output_with_deterministic_settings() -> None:
    class FakeClient:
        def __init__(self):
            self.calls = []
            self.responses = iter(["Score: 1", "1"])

        def generate_detailed(self, model, prompt, *, options, think):
            self.calls.append((model, prompt, options, think))
            return SimpleNamespace(response=next(self.responses))

    row = EvaluationRow(2, "Question", "Truth", Decimal("2"), "q1", "Comment")
    client = FakeClient()

    score = EvaluationWorker._judge(client, row, "Generated")

    assert score == Decimal("1")
    assert len(client.calls) == 2
    assert all(call[2] == {"temperature": 0, "num_predict": 16} for call in client.calls)
    assert all(call[3] is False for call in client.calls)


def test_summary_uses_all_available_points_as_accuracy_denominator() -> None:
    rows = (
        EvaluationRow(2, "Q1", "A1", Decimal("2"), "q1"),
        EvaluationRow(3, "Q2", "A2", Decimal("1"), "q2"),
    )
    outcomes = {
        2: EvaluationOutcome(rows[0], "Generated 1", Decimal("1.5")),
        3: EvaluationOutcome(rows[1], "", error="Model unavailable"),
    }

    summary = summarize_evaluation(rows, outcomes)

    assert summary.earned_points == Decimal("1.5")
    assert summary.possible_points == Decimal("3")
    assert summary.accuracy_percent == Decimal("50")
    assert summary.completed_rows == 1
    assert summary.failed_rows == 1


def test_export_preserves_input_and_adds_evaluation_columns(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Evaluation"
    worksheet.append(["Extra", "qID", "Question", "Answer", "Comment", "Points"])
    worksheet.append(["keep me", "q1", "Question one", "Truth", "Be strict", 2])
    workbook.save(source)
    evaluation = load_evaluation_workbook(source)
    row = evaluation.rows[0]
    outcome = EvaluationOutcome(row, "Generated answer", Decimal("1.5"))
    destination = tmp_path / "report.xlsx"

    parameters = EvaluationReportParameters(
        answer_model="answer-model:latest",
        answer_style="strict",
        retrieval_type="deep",
        judge_model="qwen3.6:latest",
    )
    exported = export_evaluation_report(
        evaluation,
        {row.worksheet_row: outcome},
        destination,
        parameters,
    )

    report = load_workbook(exported, data_only=False)
    try:
        output = report["Evaluation"]
        headers = {output.cell(1, column).value: column for column in range(1, output.max_column + 1)}
        assert output.cell(2, 1).value == "keep me"
        assert output.cell(2, headers[GENERATED_ANSWER_COLUMN]).value == "Generated answer"
        assert output.cell(2, headers[JUDGE_SCORE_COLUMN]).value == 1.5
        assert output.cell(2, headers[EVALUATION_ERROR_COLUMN]).value is None
        assert report.sheetnames == ["Evaluation", SUMMARY_SHEET_NAME]
        summary = report[SUMMARY_SHEET_NAME]
        values = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(2, 15)}
        assert values["Answer model"] == "answer-model:latest"
        assert values["Answer style"] == "strict"
        assert values["Retrieval type"] == "deep"
        assert values["Judge model"] == "qwen3.6:latest"
        assert values["Earned points"] == "=SUM('Evaluation'!H2:H2)"
        assert values["Possible points"] == "=SUM('Evaluation'!F2:F2)"
        assert values["Accuracy (%)"] == "=B8/B9*100"
        assert values["Completed questions"] == 1
        assert values["Failed questions"] == 0
        assert values["Pending questions"] == 0
        assert values["Total questions"] == 1
    finally:
        report.close()
