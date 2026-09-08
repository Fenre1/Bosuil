from collections import defaultdict
from pathlib import Path

import pytest

from investigation_rag.chunking.deterministic import DeterministicChunker
from investigation_rag.extractors.base import ExtractedDocument, ExtractedSection
from investigation_rag.extractors.docx_extractor import DocxExtractor
from investigation_rag.extractors.msg_extractor import _set_attachment_chain
from investigation_rag.extractors.structure import SourceBlock, html_blocks, split_ranges
from investigation_rag.extractors.txt_extractor import TxtExtractor
from investigation_rag.ingest.pipeline_state import pipeline_signatures
from investigation_rag.ingest.searchable_text import estimate_embedding_tokens, stable_embedding_text


def assert_source_alignment(document, chunks):
    blocks = {block.block_id: block for section in document.sections for block in section.blocks}
    ranges = defaultdict(list)
    for chunk in chunks:
        for span in chunk.metadata["source_spans"]:
            source = blocks[span["block_id"]]
            assert chunk.text[span["char_start"]:span["char_end"]] == source.text[span["block_start"]:span["block_end"]]
            ranges[source.block_id].append((span["block_start"], span["block_end"]))
    for block_id, block in blocks.items():
        if not block.text.strip():
            continue
        intervals = sorted(ranges[block_id])
        assert intervals[0][0] == 0
        assert intervals[-1][1] == len(block.text)
        assert all(first[1] == second[0] for first, second in zip(intervals, intervals[1:]))


def test_docx_split_keeps_exact_formatting_and_table_rows(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    document = Document()
    document.add_heading("Findings", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Repeated sentence. " * 20).bold = True
    paragraph.add_run("Final italic conclusion.").italic = True
    table = document.add_table(rows=1, cols=2)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(0, 0).text, table.cell(0, 1).text = "Control", "Owner"
    for number in range(12):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = f"Review {number}", f"Person {number}"
    document.add_heading("Actions", level=1)
    document.add_paragraph("Follow up.")
    path = tmp_path / "report.docx"
    document.save(path)
    extracted = DocxExtractor().extract(path)
    chunks = DeterministicChunker(target_chars=160, max_chars=240, overlap_chars=0).chunk_document(extracted)
    assert_source_alignment(extracted, chunks)
    assert any("<b>Repeated sentence." in chunk.metadata["display_html"] for chunk in chunks)
    assert any("<i>Final italic conclusion.</i>" in chunk.metadata["display_html"] for chunk in chunks)
    assert any(chunk.metadata.get("table_header_context") == "Control\tOwner" for chunk in chunks)
    for chunk in chunks:
        if chunk.metadata["contains_table"]:
            assert "Repeated sentence" not in chunk.text
            assert "<td " in chunk.metadata["display_html"]
        assert not ("Findings" in chunk.text and "Actions" in chunk.text)


def test_txt_keeps_whitespace_repeated_passages_and_full_input_budget(tmp_path):
    text = "    indented\tvalue\n" + "A repeated sentence. " * 100 + "\n\n" + "漢字データ。 " * 200
    path = tmp_path / "source.txt"
    path.write_bytes(text.encode())
    document = TxtExtractor().extract(path)
    document.metadata["title"] = "Very long context " * 300
    chunks = DeterministicChunker(target_chars=300, max_chars=400, max_tokens=100).chunk_document(document)
    assert_source_alignment(document, chunks)
    assert "".join(chunk.text for chunk in chunks) == text
    assert all(len(stable_embedding_text(c.text, c.metadata)) <= 400 for c in chunks)
    assert all(estimate_embedding_tokens(stable_embedding_text(c.text, c.metadata)) <= 100 for c in chunks)
    assert chunks[0].metadata["line_start"] == 1
    assert chunks[-1].metadata["line_end"] == 4


def test_docx_numbering_definitions_render_bullets_and_numbered_lists(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("One bullet", style="List Bullet")
    document.add_paragraph("First action", style="List Number")
    document.add_paragraph("Second action", style="List Number")
    path = tmp_path / "lists.docx"
    document.save(path)
    extracted = DocxExtractor().extract(path)
    blocks = extracted.sections[0].blocks
    assert [block.kind for block in blocks] == ["list_item"] * 3
    assert [block.location["list_marker"] for block in blocks] == ["•", "1.", "2."]
    chunks = DeterministicChunker().chunk_document(extracted)
    assert_source_alignment(extracted, chunks)


def test_html_email_uses_one_source_for_text_and_formatting():
    html = '<html><head><style>hidden css</style></head><body><p>Hello <b>Alice</b>.</p><script>hidden script</script><blockquote><p>Earlier message</p></blockquote><table><tr><th>Item</th><th>Cost</th></tr><tr><td>Book</td><td>12</td></tr></table><img src="https://example.com/tracking"></body></html>'
    blocks = html_blocks(html)
    document = ExtractedDocument("msg", [ExtractedSection("", "msg", blocks=blocks)])
    chunks = DeterministicChunker().chunk_document(document)
    assert_source_alignment(document, chunks)
    assert any("<b>Alice</b>" in c.metadata["display_html"] for c in chunks)
    text = " ".join(c.text for c in chunks)
    assert "hidden" not in text
    assert "tracking" not in str([c.metadata["display_html"] for c in chunks])
    assert all(not ("Hello" in c.text and "Earlier message" in c.text) for c in chunks)


def test_html_lists_keep_numbering_and_styles_stop_at_the_right_offset():
    blocks = html_blocks('<p><b>Bold</b> plain</p><ol start="3"><li>First</li><li>Second</li></ol>')
    assert [block.location["list_marker"] for block in blocks[1:]] == ["3.", "4."]
    assert not blocks[0].runs[-1]["bold"]


def test_unicode_context_leaves_room_for_source_text():
    document = ExtractedDocument("txt", [ExtractedSection("Original evidence. " * 20, "txt")],
                                 metadata={"title": "🔎" * 300})
    chunks = DeterministicChunker(target_chars=200, max_chars=400, max_tokens=32).chunk_document(document)
    assert all(estimate_embedding_tokens(stable_embedding_text(c.text, c.metadata)) <= 32 for c in chunks)
    assert all(c.metadata["embedding_token_estimate"] <= 32 for c in chunks)


def test_text_line_range_does_not_include_the_next_line(tmp_path):
    path = tmp_path / "lines.txt"
    path.write_bytes(b"First line.\nSecond line.\n")
    chunks = DeterministicChunker(target_chars=13, max_chars=26, overlap_chars=0).chunk_document(TxtExtractor().extract(path))
    assert [(c.metadata["line_start"], c.metadata["line_end"]) for c in chunks] == [(1, 1), (2, 2)]


def test_nested_attachments_get_distinct_durable_identities():
    leaf = ExtractedDocument("txt", [ExtractedSection("Leaf evidence.", "txt")])
    _set_attachment_chain(leaf, {"index": 0, "name": "same.txt"})
    nested = ExtractedDocument("msg", [ExtractedSection("Nested message.", "msg")], attachments=[leaf])
    _set_attachment_chain(nested, {"index": 1, "name": "nested.msg"})
    sibling = ExtractedDocument("txt", [ExtractedSection("Sibling evidence.", "txt")])
    _set_attachment_chain(sibling, {"index": 2, "name": "same.txt"})
    document = ExtractedDocument("msg", [ExtractedSection("Root message.", "msg")], attachments=[nested, sibling])
    chunks = DeterministicChunker().chunk_document(document)
    assert [c.metadata["source_document_id"] for c in chunks] == ["root", "1", "1/0", "2"]
    assert all(not c.metadata.get("overlap_text") for c in chunks)
    assert ExtractedDocument.from_dict(document.to_dict()).to_dict() == document.to_dict()


def test_email_signature_tracks_pdf_attachment_settings():
    initial = pipeline_signatures(".msg", DeterministicChunker(), "model")
    changed = pipeline_signatures(".msg", DeterministicChunker(pdf_soft_min_chars=100), "model")
    assert initial.chunking != changed.chunking
    assert pipeline_signatures(".txt", DeterministicChunker(), "model").chunking == pipeline_signatures(".txt", DeterministicChunker(pdf_soft_min_chars=100), "model").chunking


def test_sentence_and_word_splitting_preserves_content():
    text = "Short sentence. A longer sentence with several words. " * 20
    ranges = split_ranges(text, 72)
    assert "".join(text[a:b] for a, b in ranges) == text
    assert all(b - a <= 72 for a, b in ranges)
    assert all(text[b - 1].isspace() for a, b in ranges[:-1])


@pytest.mark.parametrize("path", list((Path(__file__).parents[1] / "test_docs").glob("*.pdf")), ids=lambda p: p.stem[:35])
def test_supplied_pdfs_have_bounded_mapped_passages(path):
    from investigation_rag.extractors.pdf_extractor import PdfExtractor
    from investigation_rag.chunking.pdf_layout import collect_layout_blocks, mark_recurring_page_furniture, PdfChunkingConfig
    document = PdfExtractor().extract(path)
    chunks = DeterministicChunker().chunk_document(document)
    assert chunks
    # Regression: font-span counting previously produced hundreds of fake tables.
    assert len(chunks) < len(document.sections) * 6
    seen = defaultdict(list)
    for chunk in chunks:
        assert estimate_embedding_tokens(stable_embedding_text(chunk.text, chunk.metadata)) <= 1400
        assert len(stable_embedding_text(chunk.text, chunk.metadata)) <= 4800
        assert chunk.metadata["source_spans"]
        for span in chunk.metadata["source_spans"]:
            assert 0 <= span["char_start"] < span["char_end"] <= len(chunk.text)
            seen[span["block_id"]].append((span["block_start"], span["block_end"]))
    for intervals in seen.values():
        intervals.sort()
        assert all(first[1] <= second[0] for first, second in zip(intervals, intervals[1:]))
    blocks = collect_layout_blocks(document)
    mark_recurring_page_furniture(blocks, PdfChunkingConfig())
    expected = {b.block_id for b in blocks if not (b.is_recurring_header or b.is_recurring_footer or b.is_page_number)}
    assert set(seen) == expected


def test_neighbor_expansion_stays_inside_an_attachment():
    from investigation_rag.retrieval.ranking import SearchResult
    from investigation_rag.retrieval.context_expansion import NeighborContextExpander
    import json
    selected = SearchResult("current", 1, 2, "mail.msg", "Evidence", "txt", 3, None, 1,
                            attachment_name="same.txt", metadata={"source_document_id": "1"})

    def row(identity, index):
        return dict(id=str(index), investigation_id=1, file_id=2, path="mail.msg", chunk_text="Neighbor",
                    source_type="txt", chunk_index=index, page_number=None, email_subject=None, email_sender=None,
                    email_recipients=None, email_date=None, attachment_name="same.txt",
                    metadata_json=json.dumps({"source_document_id": identity}))

    class DB:
        def get_neighbor_chunks(self, *args, **kwargs):
            return [row("1", 2), row("2", 4)]
    expanded = NeighborContextExpander(DB()).expand([selected], 4)
    assert [result.chunk_id for result in expanded] == ["current", "2"]
