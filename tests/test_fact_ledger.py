from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from investigation_rag.fact_ledger import (
    FACT_LEDGER_SHEET,
    MODEL_OUTPUT_SHEET,
    SOURCE_SHEET,
    FactLedgerFailure,
    FactLedgerItem,
    FactLedgerModelOutput,
    build_document_fact_ledger,
    export_fact_ledger,
    overlapping_chunk_batches,
    parse_consolidated_facts,
)


def make_chunk(index: int, *, page: int | None = None) -> dict[str, object]:
    return {
        "id": f"c{index}",
        "chunk_index": index,
        "chunk_text": f"Tekst uit chunk {index + 1}.",
        "page_number": page,
        "attachment_name": None,
    }


@pytest.mark.parametrize(
    ("count", "expected"),
    [
        (0, []),
        (1, [[0]]),
        (5, [[0, 1, 2, 3, 4]]),
        (6, [[0, 1, 2, 3, 4], [4, 5]]),
        (9, [[0, 1, 2, 3, 4], [4, 5, 6, 7, 8]]),
        (10, [[0, 1, 2, 3, 4], [4, 5, 6, 7, 8], [8, 9]]),
    ],
)
def test_batches_five_chunks_with_one_chunk_overlap(count: int, expected: list[list[int]]) -> None:
    assert overlapping_chunk_batches(list(range(count))) == expected


def test_document_ledger_processes_all_chunks_then_consolidates_with_provenance() -> None:
    class FakeClient:
        def __init__(self):
            self.prompts: list[str] = []
            self.responses = iter(
                [
                    '{"statements": ['
                    '{"statement": "De soep was heet", "source_chunk_ids": ["c0"]},'
                    '{"statement": "De soep was gemaakt met tomaten.", "source_chunk_ids": ["c4"]}'
                    "]}",
                    '{"statements": ['
                    '{"statement": "De soep was gemaakt met tomaten.", "source_chunk_ids": ["c4"]},'
                    '{"statement": "De soep werd bij het diner geserveerd.", "source_chunk_ids": ["c6"]}'
                    "]}",
                    '{"items": ['
                    '{"statement": "De soep was heet en gemaakt met tomaten.", '
                    '"source_item_ids": ["P1", "P2", "P3"]},'
                    '{"statement": "De soep werd bij het diner geserveerd.", "source_item_ids": ["P4"]}'
                    "]}",
                ]
            )

        def generate_detailed(self, model, prompt, *, options, think):
            self.prompts.append(prompt)
            assert model == "model:test"
            assert options["temperature"] == 0
            assert think is False
            return SimpleNamespace(response=next(self.responses))

    chunks = [make_chunk(index, page=index + 1) for index in range(7)]
    progress: list[str] = []
    snapshots: list[list[FactLedgerItem]] = []
    client = FakeClient()

    ledger = build_document_fact_ledger(
        client,
        "model:test",
        "rapport.pdf",
        chunks,
        progress=progress.append,
        ledger_updated=snapshots.append,
    )

    assert len(client.prompts) == 3
    assert "Chunk ID: c4" in client.prompts[0]
    assert "Chunk ID: c4" in client.prompts[1]
    assert "Chunk ID: c0" not in client.prompts[1]
    assert "Voorlopig feitenoverzicht" in client.prompts[2]
    assert ledger == [
        FactLedgerItem(
            "De soep was heet en gemaakt met tomaten.",
            ("c0", "c4"),
            ("pagina 1", "pagina 5"),
        ),
        FactLedgerItem(
            "De soep werd bij het diner geserveerd.",
            ("c6",),
            ("pagina 7",),
        ),
    ]
    assert len(snapshots) == 3
    assert progress[-1].startswith("Alle 7 chunks zijn verwerkt")


def test_consolidation_cannot_silently_drop_preliminary_items() -> None:
    preliminary = [
        FactLedgerItem("Eerste verklaring.", ("c1",)),
        FactLedgerItem("Tweede verklaring.", ("c2",)),
    ]
    response = '{"items": [{"statement": "Eerste verklaring.", "source_item_ids": ["P1"]}]}'

    with pytest.raises(ValueError, match="P2"):
        parse_consolidated_facts(
            response,
            preliminary,
            {"c1": make_chunk(1), "c2": make_chunk(2)},
        )


def test_invalid_batch_json_retries_same_chunks_then_continues_after_second_failure() -> None:
    class FakeClient:
        def __init__(self):
            self.prompts: list[str] = []
            self.responses = iter(
                [
                    "not json",
                    "still not json",
                    '{"statements": [{"statement": "Chunk zes bevat een feit.", '
                    '"source_chunk_ids": ["c5"]}]}',
                    '{"items": [{"statement": "Chunk zes bevat een feit.", '
                    '"source_item_ids": ["P1"]}]}',
                ]
            )

        def generate_detailed(self, model, prompt, *, options, think):
            self.prompts.append(prompt)
            return SimpleNamespace(response=next(self.responses))

    failures: list[FactLedgerFailure] = []
    model_outputs: list[FactLedgerModelOutput] = []
    progress: list[str] = []
    client = FakeClient()

    ledger = build_document_fact_ledger(
        client,
        "model:test",
        "rapport.pdf",
        [make_chunk(index) for index in range(6)],
        progress=progress.append,
        failure_callback=failures.append,
        model_output_callback=model_outputs.append,
    )

    assert client.prompts[0] == client.prompts[1]
    assert "Chunk ID: c5" in client.prompts[2]
    assert ledger[0].statement == "Chunk zes bevat een feit."
    assert failures == [
        FactLedgerFailure(
            stage="extractie",
            batch_number=1,
            chunk_ids=("c0", "c1", "c2", "c3", "c4"),
            chunk_numbers=(1, 2, 3, 4, 5),
            error="Het model gaf geen geldig JSON-antwoord.",
        )
    ]
    assert any("dezelfde chunks opnieuw" in message for message in progress)
    assert any("doorgaan met de volgende chunks" in message for message in progress)
    assert [output.valid_json for output in model_outputs] == [False, False, True, True]
    assert [output.attempt for output in model_outputs] == [1, 2, 1, 1]
    assert model_outputs[0].raw_output == "not json"


def test_invalid_final_json_retries_then_preserves_preliminary_ledger() -> None:
    class FakeClient:
        def __init__(self):
            self.prompts: list[str] = []
            self.responses = iter(
                [
                    '{"statements": [{"statement": "De soep was heet.", '
                    '"source_chunk_ids": ["c0"]}]}',
                    "not json",
                    "still not json",
                ]
            )

        def generate_detailed(self, model, prompt, *, options, think):
            self.prompts.append(prompt)
            return SimpleNamespace(response=next(self.responses))

    failures: list[FactLedgerFailure] = []
    client = FakeClient()

    ledger = build_document_fact_ledger(
        client,
        "model:test",
        "rapport.pdf",
        [make_chunk(0, page=3)],
        failure_callback=failures.append,
    )

    assert client.prompts[1] == client.prompts[2]
    assert ledger == [FactLedgerItem("De soep was heet.", ("c0",), ("pagina 3",))]
    assert len(failures) == 1
    assert failures[0].stage == "consolidatie"


def test_excel_export_names_document_and_has_source_path_sheet(tmp_path: Path) -> None:
    source = Path("C:/Onderzoek/Bronrapport.pdf")
    destination = tmp_path / "Bronrapport - feitenoverzicht.xlsx"
    item = FactLedgerItem("De soep was heet.", ("chunk-1",), ("pagina 4",))

    failure = FactLedgerFailure(
        stage="extractie",
        batch_number=2,
        chunk_ids=("chunk-5", "chunk-6"),
        chunk_numbers=(5, 6),
        error="Het model gaf geen geldig JSON-antwoord.",
    )
    model_output = FactLedgerModelOutput(
        stage="extractie",
        batch_number=2,
        attempt=1,
        raw_output="x" * 30_010,
        valid_json=False,
        chunk_ids=("chunk-5", "chunk-6"),
        chunk_numbers=(5, 6),
        error="Het model gaf geen geldig JSON-antwoord.",
    )

    exported = export_fact_ledger(destination, source, [item], [failure], [model_output])

    workbook = load_workbook(exported)
    try:
        assert workbook.sheetnames == [FACT_LEDGER_SHEET, SOURCE_SHEET, MODEL_OUTPUT_SHEET]
        ledger = workbook[FACT_LEDGER_SHEET]
        assert ledger["A1"].value == "Documentnaam"
        assert ledger["B1"].value == "Bronrapport.pdf"
        assert ledger["B5"].value == "De soep was heet."
        assert ledger["C5"].value == "pagina 4"
        source_sheet = workbook[SOURCE_SHEET]
        assert source_sheet["A2"].value == "Bronpad"
        assert source_sheet["B2"].value == str(source)
        assert source_sheet["A4"].value == "Mislukte verwerking"
        assert source_sheet["A6"].value == "extractie"
        assert source_sheet["B6"].value == 2
        assert source_sheet["C6"].value == "5; 6"
        assert source_sheet["D6"].value == "chunk-5; chunk-6"
        raw_output = workbook[MODEL_OUTPUT_SHEET]
        assert raw_output["A2"].value == "extractie"
        assert raw_output["B2"].value == 2
        assert raw_output["C2"].value == 1
        assert raw_output["F2"].value == "nee"
        assert raw_output["H2"].value == "1/2"
        assert raw_output["H3"].value == "2/2"
        assert raw_output["I2"].value + raw_output["I3"].value == model_output.raw_output
    finally:
        workbook.close()
